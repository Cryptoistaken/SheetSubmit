"""Generate test/Page500.xlsx — 500 Page-style rows (local-only fixture, gitignored).
Usage: python scripts/gen-page500.py
Deterministic (seeded) so UIDs are stable across runs.
"""
import random
import string

import openpyxl

N = 500
BASE_UID = 61590000000000
OUT = "test/Page500.xlsx"
rng = random.Random(42)
B32 = "ABCDEFGHIJKLMNOPQRSTUVWXYZ234567"


def rnd(n, alpha=string.ascii_letters + string.digits):
    return "".join(rng.choice(alpha) for _ in range(n))


wb = openpyxl.Workbook()
ws = wb.active
assert ws is not None
ws.title = "Sheet1"
for i in range(N):
    uid = BASE_UID + i
    cookie = (
        f"dpr=1.4375; datr={rnd(22)}; sb={rnd(22)}; m_pixel_ratio=1.4375; "
        f"wd=501x1122; c_user={uid}; fr={rnd(60)}; xs={rng.randint(1, 99)}"
        f"%3A{rnd(12)}%3A2%3A{rng.randint(1700000000, 1799999999)}%3A-1%3A-1"
    )
    key = " ".join("".join(rng.choice(B32) for _ in range(4)) for _ in range(8))
    ws.cell(row=i + 1, column=1, value=cookie)
    ws.cell(row=i + 1, column=2, value=key)
wb.save(OUT)
print(f"wrote {OUT} ({N} rows)")
